#!/usr/bin/env python3
"""Generate deterministic, non-Excel inquiry fixtures and their ground truth."""

from __future__ import annotations

import json
import math
import random
import textwrap
from email.message import EmailMessage
from pathlib import Path

from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent
SEED = 20260828
random.seed(SEED)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size=size)
    return ImageFont.load_default()


def write_text_fixture() -> None:
    text = """Subject: RE: urgent mixed steel requirement - REV 3 (please ignore Rev 1/2)

Hi Lisa,

Putting the final quantities in one mail because the WhatsApp messages got confusing.
Everything below is the FINAL request unless I explicitly say "reference only".

For the bridge package use S355JR / EN 10025-2, mill certificates 3.1:
  plate twelve by two thousand by six thousand mm ........ 80 PCS
  and 20 x 2500 x 12000 mm ................................ 24 pieces
No painting. Shot blasted Sa 2.5 and shop primer 20-25 microns. Plate tolerance EN 10029 class B.

Profiles (ASTM A36):
  equal angle L75 x 75 x 6, bars are 6 metres, need 120 pcs;
  unequal angle L100 x 75 x 8 - twelve metre bars - 60 pcs.
Bundles suitable for ocean freight.

Pipe is seamless black, bevelled ends, ASTM A106 Gr B / ASME B36.10:
  OD 168.3, wall 7.11, fixed length 11.8 m, forty-eight lengths.
  OD 219.1 x WT 8.18, random 10-12 m, quantity 32 pcs.
Hydro test certificates required.

Also add flat bar EN 10025 S275JR: 50 wide x 8 thick x 6M, 2.4 MT.
And round bar C45 / EN 10083-2: diameter 45 x 6000, 3.2 metric tons.

Commercial: CFR Callao, Peru. Shipment split is acceptable but everything must leave by 15 Oct 2026.
Payment 20% advance, balance against copy B/L. Quote in USD.

IMPORTANT correction to the old chain: the 20 mm plate is 24 PCS, not 240. The old 240 below is REFERENCE ONLY.

---- Forwarded message (obsolete Rev 1 - DO NOT USE) ----
Plate 12x2000x6000 80 pcs; Plate 20x2500x12000 240 pcs.
Incoterm FOB Shanghai. Delivery September.
---- End obsolete message ----
"""
    (ROOT / "01_forwarded_email_body.txt").write_text(text, encoding="utf-8")


def write_html_fixture() -> None:
    html = """<!doctype html>
<html><head><meta charset="utf-8"><title>RFQ Addendum</title></head>
<body style="font-family:Arial,sans-serif;color:#222">
<p><b>Subject:</b> Addendum 4 - warehouse extension (latest revision)</p>
<p>Dear team, use the blue text and the notes marked <b>FINAL</b>. Grey quoted text is the previous issue.</p>
<div style="border-left:5px solid #1677ff;padding-left:14px">
  <p><b>HOLLOW SECTIONS - EN 10219, S355J2H</b></p>
  <p>RHS 120 x 60 x 5 mm; length is shared: 12,000 mm; <b>96 pcs</b>.</p>
  <p>SHS 80 x 80 x 4 mm; same 12 m length; <b>144 pcs</b>.</p>
  <p>CHS 114.3 OD x 5.0 wall; <span style="color:#1677ff"><b>6,000 mm</b></span>; <b>75 pcs</b>.</p>
</div>
<p>All hollow sections: hot-dip galvanized after fabrication, zinc average 85 microns; ends capped; bundle labels by size.</p>
<hr>
<p><b>Merchant bars - grade S275JR unless noted</b></p>
<ul>
  <li>square solid bar 30 x 30 x 6M - 1.8 MT</li>
  <li>flat 100 wide x 12 thick, 6 metre - 2.2 MT</li>
  <li>round solid dia 60, C45, 6M - 3.6 MT</li>
  <li>angle L 65 x 65 x 5, 6M - 200 PCS</li>
</ul>
<p><b>FINAL commercial note:</b> CIF Mombasa, Kenya; delivery November 2026; payment by irrevocable LC at sight.</p>
<blockquote style="color:#888">
Previous issue - superseded: CHS was <s>12M / 150 pcs</s>; term was <s>FOB</s>. Do not quote these deleted values.
</blockquote>
<p>Regards,<br>Amara</p>
</body></html>"""
    (ROOT / "02_revision_with_markup.html").write_text(html, encoding="utf-8")


def write_voice_transcript_fixture() -> None:
    transcript = """VOICE NOTE TRANSCRIPT - buyer group / automatic transcription
Project: Tema appliance plant replenishment

[00:00] Buyer - Kojo:
Please use this voice note as the final list. Some numbers from yesterday are repeated only so you know what changed.

[00:17] Buyer - Kojo:
Cold rolled first. Standard EN 10130, grade DC01, skin-passed and lightly oiled. Coil inside diameter six-ten millimetres; maximum nine point five metric tons each.
Size zero point seven by one thousand, quantity one hundred eighty metric tons.
Next is zero point nine by twelve-fifty, two hundred forty metric tons.

[00:54] Engineer - Ama:
Painted coils are separate, yes? Pre-painted galvanized, Z one-twenty coating.

[01:02] Buyer - Kojo:
Correct. PPGI zero point four-five by one thousand, colour RAL nine-zero-zero-two, one hundred twenty MT.
Then zero point five-zero by twelve-fifty, RAL five-zero-one-zero, one hundred sixty MT.
Both regular spangle substrate and protective film on top.

[01:39] Warehouse - Mensah:
For reinforcing steel I wrote ten millimetre yesterday.

[01:45] Buyer - Kojo:
Delete ten millimetre; we do NOT need it. Final B500B to BS 4449 is twelve millimetre by twelve metres, one hundred MT; and sixteen millimetre by twelve metres, one hundred fifty MT.

[02:18] Buyer - Kojo:
Wire rod SAE one-zero-zero-eight: five point five millimetre coils, two hundred MT. Six point five millimetre coils, one hundred eighty MT. Natural finish, mill coils.

[02:43] Engineer - Ama:
Add welded mesh, eight millimetre wire, openings one-fifty by one-fifty, sheet two point four by six metres. Five hundred sheets. Standard BS 4483.

[03:10] Buyer - Kojo:
Commercial terms for every line: CFR Tema, Ghana. Ship in two lots, first half October 2026 and balance November 2026. Irrevocable LC at sight. Do not turn pieces or sheets into tons.

[03:32] System transcription note:
Low-confidence words: "six-ten" means 610 mm; "twelve-fifty" means 1250 mm. Buyer confirmed these spellings in chat.
"""
    (ROOT / "06_voice_note_transcript.txt").write_text(transcript, encoding="utf-8")


def set_docx_font(run, size: float = 11, bold: bool | None = None, color: str = "222222", italic: bool = False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.italic = italic
    if bold is not None:
        run.bold = bold


def shade_paragraph(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def add_docx_line(doc: Document, text: str, *, bullet: bool = False, color: str = "222222", bold: bool = False) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if bullet else None)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    if bullet:
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    set_docx_font(paragraph.add_run(text), bold=bold, color=color)


def write_docx_fixture() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for style_name, size, color in (("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5"), ("Heading 3", 12, "1F4D78")):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_docx_font(header.add_run("PROJECT KESTREL | MATERIAL REQUEST | REV D"), size=9, bold=True, color="6B7280")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_docx_font(footer.add_run("Fictional test fixture - values intentionally contain revisions"), size=8.5, color="6B7280")

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_docx_font(title.add_run("PROJECT KESTREL MATERIAL REQUEST"), size=23, bold=True, color="111827")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_docx_font(subtitle.add_run("Revision D - FINAL FOR QUOTATION"), size=14, bold=True, color="B42318")
    for label, value in (
        ("To", "Export Sales Team"),
        ("From", "Kestrel Fabrication JV - fictional buyer"),
        ("Date", "28 August 2026"),
        ("Rule", "Page 2 FINAL notes override matching values on page 1"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_docx_font(paragraph.add_run(f"{label}: "), bold=True)
        set_docx_font(paragraph.add_run(value))

    callout = doc.add_paragraph()
    callout.paragraph_format.space_before = Pt(12)
    callout.paragraph_format.space_after = Pt(12)
    shade_paragraph(callout, "F2F4F7")
    set_docx_font(callout.add_run("GLOBAL TERMS  "), bold=True, color="1F4D78")
    set_docx_font(callout.add_run("DDP Brno, Czech Republic; arrival by 31 Jan 2027; payment 60 days after delivery; EN 10204 3.1 certificates."))

    doc.add_heading("Frame members", level=1)
    add_docx_line(doc, "Shared grade for the first three lines: S355J2+N / EN 10025-2. Black finish, bundle by section.")
    add_docx_line(doc, "IPE 200, twelve-metre bars - 36 PCS", bullet=True)
    add_docx_line(doc, "HEB 240 x 12M - 180 PCS (OLD quantity; page 2 contains FINAL correction)", bullet=True, color="7A5A00")
    add_docx_line(doc, "UPN 120, length 6,000 mm - 90 PCS", bullet=True)

    doc.add_heading("Hollow sections", level=1)
    add_docx_line(doc, "EN 10219, grade S355J2H. Hot-dip galvanized after fabrication, average zinc 70 microns.")
    add_docx_line(doc, "RHS 160 x 80 x 5 x 12M - 44 PCS (OLD wall; see FINAL correction)", bullet=True, color="7A5A00")
    add_docx_line(doc, "SHS 100 x 100 x 5, six-metre lengths - 70 PCS", bullet=True)

    revision = doc.add_paragraph()
    revision.paragraph_format.page_break_before = True
    revision.paragraph_format.space_before = Pt(12)
    revision.paragraph_format.space_after = Pt(8)
    shade_paragraph(revision, "FCE8E6")
    set_docx_font(revision.add_run("FINAL REVISION NOTES - THESE CONTROL"), size=16, bold=True, color="B42318")

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    old = paragraph.add_run("HEB 240 quantity 180 PCS")
    set_docx_font(old, color="7A7A7A")
    old.font.strike = True
    set_docx_font(paragraph.add_run("  -> FINAL: 18 PCS."), bold=True, color="B42318")

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    old = paragraph.add_run("RHS 160 x 80 wall 5 mm")
    set_docx_font(old, color="7A7A7A")
    old.font.strike = True
    set_docx_font(paragraph.add_run("  -> FINAL: wall 6 mm; quantity remains 44 PCS; length remains 12M."), bold=True, color="B42318")

    doc.add_heading("Additional lines approved in Revision D", level=1)
    add_docx_line(doc, "Plate S355JR / EN 10025-2, 6 x 1500 x 3000 mm - 40 PCS", bullet=True)
    add_docx_line(doc, "Plate S355JR / EN 10025-2, 10 x 2000 x 6000 mm - 22 PCS", bullet=True)
    add_docx_line(doc, "Round bar C45 / EN 10083-2, diameter 70 x 6000 mm - 2.8 MT", bullet=True)
    add_docx_line(doc, "Seamless pipe ASTM A106 Gr B, OD 273.0 x wall 9.27 x random 10-12M - 14 PCS", bullet=True)
    add_docx_line(doc, "Equal angle S275JR, L120 x 120 x 10 x 12M - 28 PCS", bullet=True)

    warning = doc.add_paragraph()
    warning.paragraph_format.space_before = Pt(12)
    shade_paragraph(warning, "FFF4CE")
    set_docx_font(warning.add_run("Do not create rows for crossed-out values. Do not convert PCS to MT. Commercial terms from page 1 apply to every final line."), bold=True, color="7A5A00")

    doc.core_properties.title = "Project Kestrel Material Request - Revision D"
    doc.core_properties.author = "Fictional Buyer"
    doc.core_properties.subject = "Non-Excel inquiry extraction fixture"
    doc.save(ROOT / "07_word_revision_memo.docx")


CHAT_MESSAGES = [
    ("buyer", "Morning. Need a quick mixed-stock quote. All lengths below are final."),
    ("buyer", "1) SQ solid 25 x 25, 6M - 180 pcs\n2) SQ solid 40 x 40 x 6000 - 72 pcs"),
    ("seller", "Both square bars S235JR?"),
    ("buyer", "Yes S235JR. Black finish. Bundle by size."),
    ("buyer", "3) ROUND BAR dia 32 x 6 metres, C45 - 2.6 MT\n4) dia 50 x 6000, C45 - 4.1 MT"),
    ("buyer", "Tube part, EN10219 S355J2H:\n5) RHS 100x50x4 x 12M = 90 pcs\n6) SHS 75x75x3 x 6M = 160 pcs"),
    ("seller", "Earlier voice note said RHS 3 mm. Which wall?"),
    ("buyer", "Correction / FINAL: RHS wall is 4 mm. Ignore 3 mm."),
    ("buyer", "7) Angle L90x90x8 x 12M, S275JR, 55 pcs\n8) Flat 75 wide x 10 thick x 6M, S275JR, 1.7 MT"),
    ("buyer", "9) Seamless pipe OD 88.9 x WT 5.49 x random 5-7M, ASTM A106 Gr.B, 110 pcs"),
    ("buyer", "CFR Durban. Need shipment before 30 Sep 2026. 30% deposit / 70% copy BL."),
]


def rounded(draw: ImageDraw.ImageDraw, box: tuple[int, int, int, int], radius: int, fill: str) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill)


def write_chat_fixture() -> None:
    width, height = 1800, 3300
    image = Image.new("RGB", (width, height), "#d9e6dd")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 180), fill="#075e54")
    draw.text((80, 52), "Buyer - Steel RFQ (mobile screenshot)", font=font(48, True), fill="white")
    y = 230
    body_font = font(34)
    time_font = font(24)
    for idx, (sender, message) in enumerate(CHAT_MESSAGES, 1):
        lines: list[str] = []
        for paragraph in message.split("\n"):
            lines.extend(textwrap.wrap(paragraph, width=58) or [""])
        bubble_w = 1260
        bubble_h = 62 + len(lines) * 46 + 34
        x = 80 if sender == "buyer" else width - bubble_w - 80
        color = "#ffffff" if sender == "buyer" else "#d9fdd3"
        rounded(draw, (x, y, x + bubble_w, y + bubble_h), 28, color)
        draw.text((x + 34, y + 20), "BUYER" if sender == "buyer" else "SALES", font=font(25, True), fill="#157a6e")
        ty = y + 58
        for line in lines:
            draw.text((x + 34, ty), line, font=body_font, fill="#202124")
            ty += 46
        draw.text((x + bubble_w - 150, y + bubble_h - 32), f"09:{10 + idx:02d}", font=time_font, fill="#667781")
        y += bubble_h + 24
    image = image.filter(ImageFilter.GaussianBlur(radius=0.25))
    image.save(ROOT / "03_mobile_chat_inquiry.jpg", quality=88, optimize=True)


PDF_PAGES = [
    [
        "NORTH COAST INDUSTRIAL - RFQ 26-771 / page 1 of 3",
        "This request was assembled from site notes. Marks in the addendum on page 3 take priority.",
        "GLOBAL: CFR Guayaquil, Ecuador | Ship no later than 20 Dec 2026",
        "Payment: 15% advance; 85% against scanned B/L. Seaworthy export packing.",
        "",
        "SECTION A - HOT ROLLED COILS (ASTM A1011 CS Type B)",
        "Common: unoiled, mill edge, coil ID 508 mm, max coil weight 12 MT",
        "A-01   0.85 x 1250 mm                     300 MT",
        "A-02   1.20 x 1500 mm                     450 MT",
        "Section total shown by buyer: 750 MT",
        "",
        "SECTION B - PLATES (EN 10025-2 / S355JR)",
        "B-01   8 x 1500 x 6000 mm                 120 MT",
        "B-02   10 x 2000 x 12000 mm               200 MT   [see page 3 correction]",
        "B-03   16 x 2500 x 12000 mm                95 MT",
        "Plate condition: as rolled; EN 10029 Class B; 3.1 certificates.",
    ],
    [
        "NORTH COAST INDUSTRIAL - RFQ 26-771 / page 2 of 3",
        "CONTINUATION - quantities are not totals unless the word TOTAL is printed.",
        "",
        "SECTION C - STRUCTURAL PROFILES",
        "Common grade for C-01/C-02: S275JR, EN 10025-2, black, 6 m bundles unless stated",
        "C-01   UPN 100 x 6000 mm                  100 PCS",
        "C-02   UPN 160 x 12000 mm                  40 PCS   (12 m overrides shared 6 m)",
        "C-03   Equal angle L 50 x 50 x 5 x 6M     180 PCS  grade S235JR",
        "",
        "SECTION D - SEAMLESS PIPE / ASTM A106 GR B",
        "Black, bevelled, varnished; ASME B36.10; hydro test report required",
        "D-01   OD 114.3 x wall 6.02 x 12000 mm     80 PCS",
        "D-02   OD 168.3 x WT 7.11 x 6M             60 PCS",
        "D-03   OD 219.1 x WT 8.18 x random 10-12M  25 PCS",
        "",
        "Handwritten margin note: do not convert PCS to tons.",
    ],
    [
        "ADDENDUM - RECEIVED 28 AUG 2026 / page 3 of 3",
        "FINAL CORRECTIONS. These replace matching line references on earlier pages.",
        "",
        "B-02 quantity is 20 MT (twenty), NOT 200 MT. Dimensions remain 10 x 2000 x 12000.",
        "A-02 width is 1450 mm, NOT 1500 mm. Thickness and quantity stay 1.20 / 450 MT.",
        "D-02 length is 11.8 metres fixed, NOT 6M. Quantity stays 60 pcs.",
        "",
        "ADD TWO ITEMS:",
        "E-01 Flat bar 80 wide x 10 thick x 6000 mm, S275JR / EN10025-2, 2.5 MT",
        "E-02 Round bar diameter 55 x 6M, C45 / EN10083-2, 3.0 MT",
        "",
        "Buyer typed grand total: 1,340 MT. This total is inconsistent after corrections and",
        "mixes PCS lines; keep detail quantities as written and flag the total discrepancy.",
        "Commercial terms from page 1 remain unchanged.",
    ],
]


def add_scan_noise(image: Image.Image, page_number: int) -> Image.Image:
    px = image.load()
    for _ in range(19000):
        x = random.randrange(image.width)
        y = random.randrange(image.height)
        base = px[x, y]
        shift = random.choice((-10, -6, 6, 10))
        px[x, y] = tuple(max(0, min(255, channel + shift)) for channel in base)
    image = ImageEnhance.Contrast(image).enhance(0.93)
    image = image.filter(ImageFilter.GaussianBlur(0.35))
    angle = (-0.45, 0.32, -0.22)[page_number - 1]
    return image.rotate(angle, resample=Image.Resampling.BICUBIC, expand=False, fillcolor="#f2efe7")


def write_pdf_fixture() -> list[Path]:
    page_images: list[Path] = []
    for page_no, lines in enumerate(PDF_PAGES, 1):
        image = Image.new("RGB", (1654, 2339), "#f7f4ec")
        draw = ImageDraw.Draw(image)
        draw.rectangle((92, 92, 1562, 2247), outline="#555555", width=3)
        y = 135
        for idx, line in enumerate(lines):
            if idx == 0:
                draw.rectangle((115, y - 16, 1538, y + 64), fill="#263a4a")
                draw.text((140, y), line, font=font(32, True), fill="white")
                y += 110
                continue
            if line.startswith("SECTION") or line.startswith("ADD TWO") or line.startswith("FINAL"):
                draw.rectangle((128, y - 8, 1518, y + 55), fill="#d7e1e8")
                draw.text((145, y), line, font=font(29, True), fill="#172b3a")
                y += 82
                continue
            wrapped = textwrap.wrap(line, width=96) or [""]
            for wrapped_line in wrapped:
                draw.text((150, y), wrapped_line, font=font(27), fill="#252525")
                y += 46
            y += 12
        draw.text((1320, 2190), f"scan {page_no}/3", font=font(24), fill="#666666")
        if page_no == 3:
            draw.ellipse((1180, 1760, 1500, 2080), outline="#b23a35", width=9)
            draw.text((1235, 1880), "FINAL", font=font(44, True), fill="#b23a35")
        image = add_scan_noise(image, page_no)
        page_path = ROOT / f"tmp_scan_page_{page_no}.png"
        image.save(page_path, optimize=True)
        page_images.append(page_path)

    pdf_path = ROOT / "04_scanned_multipage_rfq.pdf"
    pdf = canvas.Canvas(str(pdf_path), pagesize=A4, invariant=1, pageCompression=1)
    page_w, page_h = A4
    for page_path in page_images:
        pdf.drawImage(str(page_path), 0, 0, width=page_w, height=page_h)
        pdf.showPage()
    pdf.save()
    return page_images


def ground_truth() -> dict:
    def item(ref: str, product: str, quantity: str, unit: str, **fields: str) -> dict:
        row = {"source_ref": ref, "product": product, "quantity": quantity, "quantity_unit": unit}
        row.update(fields)
        return row

    return {
        "schema_version": 1,
        "notes": "Expected detail rows only. Superseded and displayed total rows must not become items.",
        "cases": [
            {
                "file": "01_forwarded_email_body.txt",
                "expected_rows": 8,
                "items": [
                    item("TXT-01", "steel plate", "80", "PCS", grade="S355JR", thickness="12", width="2000", length_or_form="6000"),
                    item("TXT-02", "steel plate", "24", "PCS", grade="S355JR", thickness="20", width="2500", length_or_form="12000", remarks="Use final 24 PCS; obsolete quoted chain says 240."),
                    item("TXT-03", "equal angle", "120", "PCS", grade="ASTM A36", custom_leg1_mm="75", custom_leg2_mm="75", custom_thickness_mm="6", length_or_form="6000"),
                    item("TXT-04", "unequal angle", "60", "PCS", grade="ASTM A36", custom_leg1_mm="100", custom_leg2_mm="75", custom_thickness_mm="8", length_or_form="12000"),
                    item("TXT-05", "seamless pipe", "48", "PCS", grade="ASTM A106 Gr B", custom_diameter_mm="168.3", custom_wall_thickness_mm="7.11", length_or_form="11800"),
                    item("TXT-06", "seamless pipe", "32", "PCS", grade="ASTM A106 Gr B", custom_diameter_mm="219.1", custom_wall_thickness_mm="8.18", length_or_form="random 10000-12000"),
                    item("TXT-07", "flat bar", "2.4", "MT", grade="S275JR", custom_width_mm="50", custom_thickness_mm="8", length_or_form="6000"),
                    item("TXT-08", "round bar", "3.2", "MT", grade="C45", custom_diameter_mm="45", length_or_form="6000"),
                ],
                "shared": {"incoterm": "CFR", "port": "Callao, Peru", "delivery": "by 15 Oct 2026"},
            },
            {
                "file": "02_revision_with_markup.html",
                "expected_rows": 7,
                "items": [
                    item("HTML-01", "rectangular hollow section", "96", "PCS", grade="S355J2H", custom_width_mm="120", custom_height_mm="60", custom_wall_thickness_mm="5", length_or_form="12000"),
                    item("HTML-02", "square hollow section", "144", "PCS", grade="S355J2H", custom_width_mm="80", custom_height_mm="80", custom_wall_thickness_mm="4", length_or_form="12000"),
                    item("HTML-03", "circular hollow section", "75", "PCS", grade="S355J2H", custom_diameter_mm="114.3", custom_wall_thickness_mm="5.0", length_or_form="6000"),
                    item("HTML-04", "square solid bar", "1.8", "MT", grade="S275JR", custom_width_mm="30", custom_height_mm="30", length_or_form="6000"),
                    item("HTML-05", "flat bar", "2.2", "MT", grade="S275JR", custom_width_mm="100", custom_thickness_mm="12", length_or_form="6000"),
                    item("HTML-06", "round solid bar", "3.6", "MT", grade="C45", custom_diameter_mm="60", length_or_form="6000"),
                    item("HTML-07", "equal angle", "200", "PCS", grade="S275JR", custom_leg1_mm="65", custom_leg2_mm="65", custom_thickness_mm="5", length_or_form="6000"),
                ],
                "shared": {"incoterm": "CIF", "port": "Mombasa, Kenya", "delivery": "November 2026"},
            },
            {
                "file": "03_mobile_chat_inquiry.jpg",
                "expected_rows": 9,
                "items": [
                    item("CHAT-01", "square solid bar", "180", "PCS", grade="S235JR", custom_width_mm="25", custom_height_mm="25", length_or_form="6000"),
                    item("CHAT-02", "square solid bar", "72", "PCS", grade="S235JR", custom_width_mm="40", custom_height_mm="40", length_or_form="6000"),
                    item("CHAT-03", "round bar", "2.6", "MT", grade="C45", custom_diameter_mm="32", length_or_form="6000"),
                    item("CHAT-04", "round bar", "4.1", "MT", grade="C45", custom_diameter_mm="50", length_or_form="6000"),
                    item("CHAT-05", "rectangular hollow section", "90", "PCS", grade="S355J2H", custom_width_mm="100", custom_height_mm="50", custom_wall_thickness_mm="4", length_or_form="12000", remarks="Final correction is 4 mm wall; ignore earlier 3 mm voice note."),
                    item("CHAT-06", "square hollow section", "160", "PCS", grade="S355J2H", custom_width_mm="75", custom_height_mm="75", custom_wall_thickness_mm="3", length_or_form="6000"),
                    item("CHAT-07", "equal angle", "55", "PCS", grade="S275JR", custom_leg1_mm="90", custom_leg2_mm="90", custom_thickness_mm="8", length_or_form="12000"),
                    item("CHAT-08", "flat bar", "1.7", "MT", grade="S275JR", custom_width_mm="75", custom_thickness_mm="10", length_or_form="6000"),
                    item("CHAT-09", "seamless pipe", "110", "PCS", grade="ASTM A106 Gr B", custom_diameter_mm="88.9", custom_wall_thickness_mm="5.49", length_or_form="random 5000-7000"),
                ],
                "shared": {"incoterm": "CFR", "port": "Durban", "delivery": "before 30 Sep 2026"},
                "must_resolve": ["RHS wall is 4 mm, not the superseded 3 mm", "6M means numeric length 6000 where the selected template length field is numeric"],
            },
            {
                "file": "04_scanned_multipage_rfq.pdf",
                "expected_rows": 13,
                "items": [
                    item("A-01", "hot rolled coil", "300", "MT", material_standard="ASTM A1011 CS Type B", thickness="0.85", width="1250", coil_id="508", coil_weight="12"),
                    item("A-02", "hot rolled coil", "450", "MT", material_standard="ASTM A1011 CS Type B", thickness="1.20", width="1450", coil_id="508", coil_weight="12", remarks="Page 3 changes width from 1500 to 1450 mm."),
                    item("B-01", "steel plate", "120", "MT", grade="S355JR", thickness="8", width="1500", length_or_form="6000"),
                    item("B-02", "steel plate", "20", "MT", grade="S355JR", thickness="10", width="2000", length_or_form="12000", remarks="Page 3 changes quantity from 200 to 20 MT."),
                    item("B-03", "steel plate", "95", "MT", grade="S355JR", thickness="16", width="2500", length_or_form="12000"),
                    item("C-01", "UPN channel 100", "100", "PCS", grade="S275JR", length_or_form="6000"),
                    item("C-02", "UPN channel 160", "40", "PCS", grade="S275JR", length_or_form="12000"),
                    item("C-03", "equal angle", "180", "PCS", grade="S235JR", custom_leg1_mm="50", custom_leg2_mm="50", custom_thickness_mm="5", length_or_form="6000"),
                    item("D-01", "seamless pipe", "80", "PCS", grade="ASTM A106 Gr B", custom_diameter_mm="114.3", custom_wall_thickness_mm="6.02", length_or_form="12000"),
                    item("D-02", "seamless pipe", "60", "PCS", grade="ASTM A106 Gr B", custom_diameter_mm="168.3", custom_wall_thickness_mm="7.11", length_or_form="11800", remarks="Page 3 changes length from 6M to fixed 11.8 metres."),
                    item("D-03", "seamless pipe", "25", "PCS", grade="ASTM A106 Gr B", custom_diameter_mm="219.1", custom_wall_thickness_mm="8.18", length_or_form="random 10000-12000"),
                    item("E-01", "flat bar", "2.5", "MT", grade="S275JR", custom_width_mm="80", custom_thickness_mm="10", length_or_form="6000"),
                    item("E-02", "round bar", "3.0", "MT", grade="C45", custom_diameter_mm="55", length_or_form="6000"),
                ],
                "shared": {"incoterm": "CFR", "port": "Guayaquil, Ecuador", "delivery": "no later than 20 Dec 2026"},
                "must_resolve": ["A-02 width 1450", "B-02 quantity 20", "D-02 length 11800", "ignore section totals as rows", "flag inconsistent 1,340 MT grand total"],
            },
            {
                "file": "06_voice_note_transcript.txt",
                "expected_rows": 9,
                "items": [
                    item("VOICE-01", "cold rolled coil", "180", "MT", material_standard="EN 10130", grade="DC01", thickness="0.7", width="1000", coil_id="610", coil_weight="9.5"),
                    item("VOICE-02", "cold rolled coil", "240", "MT", material_standard="EN 10130", grade="DC01", thickness="0.9", width="1250", coil_id="610", coil_weight="9.5"),
                    item("VOICE-03", "pre-painted galvanized coil", "120", "MT", coating="Z120; RAL 9002", thickness="0.45", width="1000"),
                    item("VOICE-04", "pre-painted galvanized coil", "160", "MT", coating="Z120; RAL 5010", thickness="0.50", width="1250"),
                    item("VOICE-05", "reinforcing bar", "100", "MT", material_standard="BS 4449", grade="B500B", custom_diameter_mm="12", length_or_form="12000"),
                    item("VOICE-06", "reinforcing bar", "150", "MT", material_standard="BS 4449", grade="B500B", custom_diameter_mm="16", length_or_form="12000", remarks="Do not create the superseded 10 mm line."),
                    item("VOICE-07", "wire rod", "200", "MT", grade="SAE 1008", custom_diameter_mm="5.5", length_or_form="coil"),
                    item("VOICE-08", "wire rod", "180", "MT", grade="SAE 1008", custom_diameter_mm="6.5", length_or_form="coil"),
                    item("VOICE-09", "welded mesh", "500", "SHEETS", material_standard="BS 4483", custom_diameter_mm="8", width="2400", length_or_form="6000", remarks="Opening 150 x 150 mm."),
                ],
                "shared": {"incoterm": "CFR", "port": "Tema, Ghana", "delivery": "two lots: first half Oct 2026 and balance Nov 2026"},
                "must_resolve": ["six-ten means coil ID 610", "twelve-fifty means width 1250", "do not create the obsolete 10 mm rebar line"],
            },
            {
                "file": "07_word_revision_memo.docx",
                "expected_rows": 10,
                "items": [
                    item("WORD-01", "IPE 200", "36", "PCS", grade="S355J2+N", length_or_form="12000"),
                    item("WORD-02", "HEB 240", "18", "PCS", grade="S355J2+N", length_or_form="12000", remarks="Revision D replaces old 180 PCS with 18 PCS."),
                    item("WORD-03", "UPN 120", "90", "PCS", grade="S355J2+N", length_or_form="6000"),
                    item("WORD-04", "rectangular hollow section", "44", "PCS", grade="S355J2H", custom_width_mm="160", custom_height_mm="80", custom_wall_thickness_mm="6", length_or_form="12000", remarks="Revision D replaces old 5 mm wall with 6 mm."),
                    item("WORD-05", "square hollow section", "70", "PCS", grade="S355J2H", custom_width_mm="100", custom_height_mm="100", custom_wall_thickness_mm="5", length_or_form="6000"),
                    item("WORD-06", "steel plate", "40", "PCS", grade="S355JR", thickness="6", width="1500", length_or_form="3000"),
                    item("WORD-07", "steel plate", "22", "PCS", grade="S355JR", thickness="10", width="2000", length_or_form="6000"),
                    item("WORD-08", "round bar", "2.8", "MT", grade="C45", custom_diameter_mm="70", length_or_form="6000"),
                    item("WORD-09", "seamless pipe", "14", "PCS", grade="ASTM A106 Gr B", custom_diameter_mm="273.0", custom_wall_thickness_mm="9.27", length_or_form="random 10000-12000"),
                    item("WORD-10", "equal angle", "28", "PCS", grade="S275JR", custom_leg1_mm="120", custom_leg2_mm="120", custom_thickness_mm="10", length_or_form="12000"),
                ],
                "shared": {"incoterm": "DDP", "port": "Brno, Czech Republic", "delivery": "arrival by 31 Jan 2027"},
                "must_resolve": ["HEB quantity 18, not crossed-out 180", "RHS wall 6, not crossed-out 5", "page 1 commercial terms apply to page 2 additions"],
            },
            {
                "file": "05_complete_mail_bundle.eml",
                "expected_attachments": ["02_revision_with_markup.html", "03_mobile_chat_inquiry.jpg", "04_scanned_multipage_rfq.pdf"],
                "purpose": "mail ingestion fixture; extraction is run separately per selected body or attachment",
            },
        ],
    }


def write_readme() -> None:
    readme = """# 非 Excel 复杂询盘测试集

这组样例专门验证邮件转询盘在非结构化输入上的稳定性。所有公司、联系人和订单均为虚构。

## 文件

- `01_forwarded_email_body.txt`：正文转发链，包含过期数量和最终更正。
- `02_revision_with_markup.html`：HTML 修订稿，包含删除线、共享条件和跨段继承。
- `03_mobile_chat_inquiry.jpg`：手机聊天截图，规格分散在多条消息中并有口径纠正。
- `04_scanned_multipage_rfq.pdf`：三页纯扫描 PDF，续页、共享条件、后页更正和错误总计并存。
- `05_complete_mail_bundle.eml`：完整 MIME 邮件，正文加三个非 Excel 附件。
- `06_voice_note_transcript.txt`：多人语音留言转写，数字以口语表达且包含废弃规格。
- `07_word_revision_memo.docx`：两页 Word 修订稿，删除线旧值和后页 FINAL 更正并存。
- `expected.json`：机器可读的期望行数、关键字段和必须处理的歧义。

## 验收原则

1. 明细行数必须等于 `expected_rows`，TOTAL、旧版本和删除线内容不能生成额外明细。
2. 后发且明确标记为 FINAL/correction 的值覆盖旧值，但旧值应在备注中留下可审计说明。
3. `6M`、`12 m` 等长度在数值长度列中应转换为 `6000`、`12000`；随机长度保留范围含义。
4. PCS 与 MT 不互相换算；汇总数量冲突不能靠模型猜测修正。
5. PDF 每页都必须识别，第三页的更正必须回写到第一页/第二页对应行。

运行 `generate_non_excel_cases.py` 可重复生成二进制图片、PDF 和 EML。生成结果固定使用随机种子 20260828。

## 真实模型回归

测试默认跳过，只有显式提供真实模型配置时才消耗 token：

```bash
ERP_LIVE_OPENAI=1 go test ./internal/adapter/openai \\
  -run TestLiveNonExcelInquiryFixtures -count=1 -v
```

2026-08-28 基线结果：转发 TXT 8/8、HTML 7/7、扫描 PDF 13/13、语音转写 TXT 9/9、Word 修订稿 10/10；聊天截图被图片审计拒绝。拒绝原因是审计把普通共享备注也要求逐行完全复制，并非已经确认漏行。该样例故意保留，用于推动审计规则区分“必须逐行继承的业务字段”和“只需保留一次的说明”。
"""
    (ROOT / "README.md").write_text(readme, encoding="utf-8")


def write_eml_fixture(pdf_path: Path) -> None:
    msg = EmailMessage()
    msg["From"] = "purchasing@fictional-northcoast.example"
    msg["To"] = "sales@example.test"
    msg["Subject"] = "RFQ 26-771 plus warehouse addendum - non Excel package"
    msg["Date"] = "Fri, 28 Aug 2026 09:45:00 +0800"
    msg["Message-ID"] = "<rfq-26-771@example.test>"
    msg.set_content(
        "Please treat the attached PDF, marked-up HTML and mobile chat screenshot as one inquiry package.\n"
        "Where values conflict, explicit FINAL corrections take priority. Do not convert PCS into MT.\n"
    )
    html = (ROOT / "02_revision_with_markup.html").read_text(encoding="utf-8")
    msg.add_alternative("<p>Please review the three attached non-Excel sources.</p>" + html, subtype="html")
    msg.add_attachment((ROOT / "02_revision_with_markup.html").read_bytes(), maintype="text", subtype="html", filename="02_revision_with_markup.html")
    msg.add_attachment((ROOT / "03_mobile_chat_inquiry.jpg").read_bytes(), maintype="image", subtype="jpeg", filename="03_mobile_chat_inquiry.jpg")
    msg.add_attachment(pdf_path.read_bytes(), maintype="application", subtype="pdf", filename="04_scanned_multipage_rfq.pdf")
    msg.set_boundary("erp-go-non-excel-mixed-20260828")
    for part in msg.iter_parts():
        if part.get_content_type() == "multipart/alternative":
            part.set_boundary("erp-go-non-excel-alternative-20260828")
    (ROOT / "05_complete_mail_bundle.eml").write_bytes(msg.as_bytes())


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    write_text_fixture()
    write_html_fixture()
    write_voice_transcript_fixture()
    write_docx_fixture()
    write_chat_fixture()
    page_images = write_pdf_fixture()
    truth = ground_truth()
    (ROOT / "expected.json").write_text(json.dumps(truth, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_readme()
    write_eml_fixture(ROOT / "04_scanned_multipage_rfq.pdf")
    for page_image in page_images:
        page_image.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
